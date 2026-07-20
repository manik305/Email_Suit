"""
dnc.py  –  /api/v1/dnc
~~~~~~~~~~~~~~~~~~~~~~~
Do Not Contact (DNC) List management.

Routes:
  GET    /                          – list all DNC entries (filterable by project_id)
  POST   /classify                  – classify a response → add to DNC, update recipient, stop follow-ups
  DELETE /{dnc_id}                  – remove from DNC (un-blacklist)
  GET    /check/{email}             – check if an email is on the DNC list
  GET    /export/{campaign_id}      – export classified responses for a campaign as Excel
"""

import logging
import io
from typing import List, Optional

from fastapi import APIRouter, HTTPException, Query, Depends
from fastapi.responses import StreamingResponse

from .. import models, schemas
from .auth import get_current_user_email

logger = logging.getLogger(__name__)
router = APIRouter()


# ─── List DNC entries ─────────────────────────────────────────────────────────

@router.get("/", response_model=List[schemas.DncEntryOut])
async def list_dnc_entries(
    project_id: Optional[str] = Query(None),
    current_user_email: str = Depends(get_current_user_email),
):
    """List all DNC entries, optionally filtered by project_id."""
    if project_id:
        entries = await models.DncEntry.find(project_id=project_id).sort("-created_at").to_list()
    else:
        entries = await models.DncEntry.find_all().sort("-created_at").to_list()
    return entries


# ─── Classify a response ─────────────────────────────────────────────────────

@router.post("/classify", response_model=schemas.DncEntryOut, status_code=201)
async def classify_response(
    payload: schemas.DncEntryCreate,
    current_user_email: str = Depends(get_current_user_email),
):
    """
    Classify an inbox response. This:
    1. Creates/updates a DNC entry for the email (project-scoped)
    2. Updates matching recipient(s) in the campaign with response_category
    3. Cancels any pending follow-ups for that recipient
    """
    valid_reasons = {"lead", "hot", "cold", "negative", "bounce", "ooo", "automatic", "not_bounce"}
    if payload.reason not in valid_reasons:
        raise HTTPException(
            status_code=400,
            detail=f"Invalid reason '{payload.reason}'. Must be one of: {', '.join(valid_reasons)}"
        )

    # Resolve project_id from campaign if not provided
    project_id = payload.project_id
    if not project_id and payload.source_campaign_id:
        campaign = await models.Campaign.get(payload.source_campaign_id)
        if campaign:
            project_id = campaign.project_id

    if not project_id:
        raise HTTPException(status_code=400, detail="project_id is required (or provide a valid source_campaign_id)")

    # Check if email already exists in DNC for this project
    existing = await models.DncEntry.find_one(email=payload.email, project_id=project_id)
    if existing:
        # Update the existing entry with new classification
        await existing.update({"$set": {
            "reason": payload.reason,
            "source_campaign_id": payload.source_campaign_id,
            "notes": payload.notes,
            "classified_by": current_user_email,
        }})
        # Reload
        existing = await models.DncEntry.get(existing.id)
    else:
        # Create new DNC entry
        existing = models.DncEntry(
            email=payload.email,
            reason=payload.reason,
            source_campaign_id=payload.source_campaign_id,
            project_id=project_id,
            notes=payload.notes,
            classified_by=current_user_email,
        )
        await existing.insert()

    # Update all recipients with this email in this campaign
    if payload.source_campaign_id:
        recipients = await models.Recipient.find(
            email=payload.email,
            campaign_id=payload.source_campaign_id,
        ).to_list()
        from app.neo4j_sync import sync_lead_response
        for r in recipients:
            await r.update({"$set": {
                "response_category": payload.reason,
                "next_follow_up_at": None,  # Cancel follow-ups
                "status": "replied" if r.status in ("sent", "pending") else r.status,
            }})
            try:
                await sync_lead_response(
                    recipient_id=r.id,
                    category=payload.reason,
                    response_text=payload.notes or ""
                )
            except Exception as sync_err:
                logger.error("Failed to sync lead classification to Neo4j: %s", sync_err)

    # Also update recipients across ALL campaigns in this project that have the same email
    from app.database import db_pool
    if db_pool:
        try:
            async with db_pool.acquire() as conn:
                # Get all campaign IDs in this project
                campaign_ids = await conn.fetch(
                    "SELECT id FROM public.campaigns WHERE project_id = $1",
                    project_id,
                )
                cids = [str(row["id"]) for row in campaign_ids]
                if cids:
                    # Cancel follow-ups for this email across all campaigns in the project
                    await conn.execute(
                        """
                        UPDATE public.recipients 
                        SET next_follow_up_at = NULL, 
                            response_category = $1
                        WHERE email = $2 
                        AND campaign_id = ANY($3::uuid[])
                        AND (response_category IS NULL OR response_category != $1)
                        """,
                        payload.reason,
                        payload.email,
                        cids,
                    )
        except Exception as e:
            logger.error("Failed to propagate DNC across project campaigns: %s", e)

    logger.info(
        "DNC classified: %s as %s in project %s by %s",
        payload.email, payload.reason, project_id, current_user_email
    )
    return existing


# ─── Remove from DNC ─────────────────────────────────────────────────────────

@router.delete("/{dnc_id}", status_code=204)
async def remove_dnc_entry(
    dnc_id: str,
    current_user_email: str = Depends(get_current_user_email),
):
    """Remove an email from the DNC list (un-blacklist)."""
    entry = await models.DncEntry.get(dnc_id)
    if not entry:
        raise HTTPException(status_code=404, detail="DNC entry not found")
    
    # Clear response_category on matching recipients
    from app.database import db_pool
    if db_pool and entry.project_id:
        try:
            async with db_pool.acquire() as conn:
                campaign_ids = await conn.fetch(
                    "SELECT id FROM public.campaigns WHERE project_id = $1",
                    entry.project_id,
                )
                cids = [str(row["id"]) for row in campaign_ids]
                if cids:
                    await conn.execute(
                        """
                        UPDATE public.recipients 
                        SET response_category = NULL
                        WHERE email = $1 
                        AND campaign_id = ANY($2::uuid[])
                        """,
                        entry.email,
                        cids,
                    )
        except Exception as e:
            logger.error("Failed to clear DNC category on recipients: %s", e)

    await entry.delete()
    logger.info("DNC entry removed: %s by %s", entry.email, current_user_email)


# ─── Check DNC status ────────────────────────────────────────────────────────

@router.get("/check/{email}")
async def check_dnc(
    email: str,
    project_id: Optional[str] = Query(None),
    current_user_email: str = Depends(get_current_user_email),
):
    """Check if an email is on the DNC list for a given project."""
    if project_id:
        entry = await models.DncEntry.find_one(email=email, project_id=project_id)
    else:
        entry = await models.DncEntry.find_one(email=email)
    
    if entry:
        return {
            "is_dnc": True,
            "reason": entry.reason,
            "classified_by": entry.classified_by,
            "created_at": entry.created_at.isoformat() if entry.created_at else None,
            "notes": entry.notes,
        }
    return {"is_dnc": False}


# ─── Export DNC responses as Excel ────────────────────────────────────────────

@router.get("/export/{campaign_id}")
async def export_dnc_responses(
    campaign_id: str,
    current_user_email: str = Depends(get_current_user_email),
):
    """Export all classified responses for a campaign as an Excel (.xlsx) file."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    campaign = await models.Campaign.get(campaign_id)
    if not campaign:
        raise HTTPException(status_code=404, detail="Campaign not found")

    # Get all recipients with a response_category
    recipients = await models.Recipient.find(campaign_id=campaign_id).to_list()
    classified = [r for r in recipients if r.response_category]

    # Create workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Classified Responses"

    # Header styling
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    headers = [
        "Email", "Name", "First Name", "Last Name", "Company",
        "Designation", "Department", "Industry", "Region",
        "Classification", "Status", "Created At"
    ]

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border

    # Category color mapping
    category_fills = {
        "lead": PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid"),
        "hot": PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid"),
        "cold": PatternFill(start_color="DBEAFE", end_color="DBEAFE", fill_type="solid"),
        "negative": PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid"),
        "bounce": PatternFill(start_color="F3F4F6", end_color="F3F4F6", fill_type="solid"),
    }

    for row_idx, r in enumerate(classified, 2):
        values = [
            r.email,
            r.name or "",
            r.first_name or "",
            r.last_name or "",
            r.company_name or "",
            r.title or "",
            r.department or "",
            r.industry or "",
            r.region or "",
            (r.response_category or "").upper(),
            r.status,
            r.created_at.strftime("%Y-%m-%d %H:%M") if r.created_at else "",
        ]
        fill = category_fills.get(r.response_category, None)
        for col, val in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col, value=val)
            cell.border = thin_border
            if fill:
                cell.fill = fill

    # Auto-width columns
    for col in range(1, len(headers) + 1):
        max_len = max(
            len(str(ws.cell(row=r, column=col).value or ""))
            for r in range(1, ws.max_row + 1)
        )
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = min(max_len + 4, 40)

    # Summary sheet
    ws2 = wb.create_sheet("Summary")
    ws2.cell(row=1, column=1, value="Campaign").font = Font(bold=True)
    ws2.cell(row=1, column=2, value=campaign.name)
    ws2.cell(row=2, column=1, value="Total Classified").font = Font(bold=True)
    ws2.cell(row=2, column=2, value=len(classified))

    categories = {"lead": 0, "hot": 0, "cold": 0, "negative": 0, "bounce": 0}
    for r in classified:
        if r.response_category in categories:
            categories[r.response_category] += 1

    row = 4
    ws2.cell(row=row, column=1, value="Category").font = Font(bold=True)
    ws2.cell(row=row, column=2, value="Count").font = Font(bold=True)
    for cat, count in categories.items():
        row += 1
        ws2.cell(row=row, column=1, value=cat.upper())
        ws2.cell(row=row, column=2, value=count)

    # Write to buffer
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    safe_name = campaign.name.replace(" ", "_").replace("/", "-")[:50]
    filename = f"DNC_Responses_{safe_name}.xlsx"

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )
